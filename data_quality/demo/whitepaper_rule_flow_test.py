"""
Whitepaper rule flow verification.

This script reads the data quality rule list from the whitepaper, finds the
corresponding rule-library entry, and runs each rule through the workflow:
parameter configuration, script preview, trial run, task execution, issue
generation, dashboard generation and report rendering.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence

from docx import Document

from quality_workflow_smoke_test import DATA_QUALITY, ROOT, install_runtime_stubs, load_module


RULE_ID_PATTERN = re.compile(r"^[A-Z]{1,2}-[A-Z][0-9]{2}$")
DIMENSION_BY_TABLE = {
    2: ("normativity", "规范性"),
    3: ("normativity", "规范性"),
    4: ("normativity", "规范性"),
    5: ("normativity", "规范性"),
    6: ("normativity", "规范性"),
    7: ("completeness", "完整性"),
    8: ("accuracy", "准确性"),
    9: ("consistency", "一致性"),
    10: ("timeliness", "时效性"),
    11: ("accessibility", "可访问性"),
}


def _load_modules() -> Dict[str, Any]:
    install_runtime_stubs()
    return {
        "rule_library": load_module(
            "metadata.data_quality.rules.rule_library",
            "rules/rule_library.py",
        ),
        "workflow": load_module(
            "metadata.data_quality.workflow.quality_assessment_workflow",
            "workflow/quality_assessment_workflow.py",
        ),
    }


def _find_whitepaper(path: str = "") -> Path:
    if path:
        candidate = Path(path)
        if candidate.exists():
            return candidate
        raise FileNotFoundError(f"Whitepaper not found: {candidate}")
    docx_files = sorted(ROOT.glob("*.docx"), key=lambda item: item.stat().st_size, reverse=True)
    if not docx_files:
        raise FileNotFoundError("No docx files found in workspace root.")
    return docx_files[0]


def _extract_whitepaper_rules(path: Path) -> List[Dict[str, Any]]:
    doc = Document(str(path))
    rows: List[Dict[str, Any]] = []
    for table_index in range(2, 12):
        dimension, dimension_zh = DIMENSION_BY_TABLE[table_index]
        table = doc.tables[table_index]
        for row_index, row in enumerate(table.rows[1:], start=1):
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if len(cells) < 2:
                continue
            whitepaper_rule_id = cells[0].strip()
            if not RULE_ID_PATTERN.match(whitepaper_rule_id):
                continue
            rows.append(
                {
                    "whitepaper_rule_id": whitepaper_rule_id,
                    "display_name": cells[1].strip(),
                    "definition": cells[2].strip() if len(cells) > 2 else "",
                    "scenario": cells[3].strip() if len(cells) > 3 else "",
                    "table_index": table_index,
                    "table_row": row_index,
                    "dimension": dimension,
                    "dimension_zh": dimension_zh,
                }
            )
    return rows


def _whitepaper_rule_key(row: Dict[str, Any]) -> str:
    return f"T{row['table_index']}-R{row['table_row']}-{row['whitepaper_rule_id']}"


def _find_rule_for_whitepaper_row(
    rules: Sequence[Any],
    row: Dict[str, Any],
) -> Any:
    original_id = row["whitepaper_rule_id"]
    display_name = row["display_name"]

    exact_candidates = []
    fallback_candidates = []
    for rule in rules:
        params = getattr(rule, "parameters", {}) or {}
        tags = set(getattr(rule, "tags", []) or [])
        has_whitepaper_id = (
            params.get("whitepaper_rule_id") == original_id
            or rule.rule_id == original_id
            or original_id in tags
        )
        if not has_whitepaper_id:
            continue
        if rule.display_name == display_name:
            exact_candidates.append(rule)
        else:
            fallback_candidates.append(rule)

    if exact_candidates:
        added = [rule for rule in exact_candidates if "白皮书补齐" in (rule.tags or [])]
        return added[0] if added else exact_candidates[0]
    if fallback_candidates:
        return fallback_candidates[0]
    raise AssertionError(
        f"Missing rule-library entry for whitepaper rule {original_id}: {display_name}"
    )


def _slug(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z_]+", "_", value).strip("_").lower() or "rule"


def _target_column(row: Dict[str, Any], rule: Any) -> str:
    entity_type = rule.applicability.entity_type.value
    if entity_type != "COLUMN":
        return ""
    text = f"{row['whitepaper_rule_id']} {row['display_name']} {row['scenario']}"
    candidates = [
        ("phone", ["手机号", "联系方式"]),
        ("email", ["邮箱"]),
        ("id_card", ["身份证", "证件"]),
        ("code", ["编码", "代码"]),
        ("event_time", ["时间", "日期"]),
        ("amount", ["金额", "价格", "数值"]),
        ("status", ["状态", "枚举", "布尔"]),
        ("unit", ["单位"]),
        ("name", ["命名", "名称"]),
        ("rate", ["比例", "完成率", "百分比"]),
    ]
    for column, keys in candidates:
        if any(key in text for key in keys):
            return column
    return "value"


def _sample_rows(rule_id: str, target_column: str) -> List[Dict[str, Any]]:
    failing_row = {
        "id": "row-001",
        "value": "invalid",
        "code": "bad-code",
        "amount": "-1",
        "status": "未知",
        "event_time": "2099-01-01 00:00:00",
        "__fail_rules": [rule_id],
        "__fail_reason": "白皮书规则测试样例标记未通过",
    }
    passing_row = {
        "id": "row-002",
        "value": "valid",
        "code": "STD001",
        "amount": "100.00",
        "status": "有效",
        "event_time": "2026-06-19 00:00:00",
    }
    if target_column:
        failing_row.setdefault(target_column, failing_row.get("value"))
        passing_row.setdefault(target_column, passing_row.get("value"))
    return [failing_row, passing_row]


def _assert(condition: bool, message: str, detail: Any = None) -> None:
    if not condition:
        raise AssertionError(f"{message}: {detail}")


def _test_rule(modules: Dict[str, Any], rule: Any, row: Dict[str, Any]) -> Dict[str, Any]:
    workflow_mod = modules["workflow"]
    rule_library_mod = modules["rule_library"]
    library = rule_library_mod.RuleLibrary(load_builtin=True)
    workflow = workflow_mod.QualityAssessmentWorkflow(rule_library=library)

    target_column = _target_column(row, rule)
    table_name = "whitepaper_rule_fixture"
    scope = workflow_mod.DataScope(
        data_source="whitepaper_fixture",
        database="dq",
        schema="test",
        table_fqn=f"whitepaper_fixture.dq.test.{table_name}",
        table_name=table_name,
        fields=[
            "id",
            "value",
            "code",
            "amount",
            "status",
            "event_time",
            target_column or "table_level_marker",
        ],
        business_domain="白皮书数据质量规则验证",
        batch_id="whitepaper_rule_flow",
        row_count=2,
    )
    setting = workflow.configure_rule_parameters(
        rule.rule_id,
        scope,
        target_column=target_column,
        parameter_overrides={"invalid_condition": "1 = 1"},
    )
    preview = workflow.preview_rule_script(setting.setting_id).to_dict()
    _assert(rule.rule_id in preview["rule_id"], "script preview rule id mismatch", preview)

    rows = _sample_rows(rule.rule_id, target_column)
    trial = workflow.trial_run(setting.setting_id, rows)
    _assert(trial.total_rows == 2, "trial total rows mismatch", trial.to_dict())
    _assert(trial.failure_count == 1, "trial did not detect the fixture failure", trial.to_dict())
    _assert(not trial.passed, "trial unexpectedly passed", trial.to_dict())

    task = workflow.create_task(
        task_name=f"白皮书规则验证-{rule.rule_id}",
        scope=scope,
        rule_setting_ids=[setting.setting_id],
        schedule="manual",
        scan_mode="full",
    )
    run = workflow.execute_task(task.task_id, sample_rows=rows, batch_id="whitepaper_rule_flow")
    _assert(run.total_rules == 1, "task did not execute exactly one rule", run.to_dict())
    _assert(run.failed_rules == 1, "task did not fail the marked rule", run.to_dict())
    _assert(len(run.issue_ids) == 1, "task did not generate one issue", run.to_dict())

    issues = workflow.query_issues(include_archived=True)
    _assert(len(issues) == 1, "issue query mismatch", [issue.to_dict() for issue in issues])
    lineage = workflow.analyze_issue_lineage(run.issue_ids[0])
    _assert(lineage["issue_id"] == run.issue_ids[0], "lineage issue mismatch", lineage)

    dashboard = workflow.build_dashboard(run_ids=[run.run_id]).to_dict()
    _assert(
        any(item["dimension"] == rule.dimension.value for item in dashboard["dimension_scores"]),
        "dashboard missing rule dimension",
        dashboard,
    )
    report = workflow.generate_workflow_report(run_id=run.run_id)
    _assert("质量评价流程执行报告" in report, "markdown report missing title", report[:120])

    return {
        "whitepaper_key": _whitepaper_rule_key(row),
        "whitepaper_rule_id": row["whitepaper_rule_id"],
        "rule_id": rule.rule_id,
        "display_name": rule.display_name,
        "dimension": rule.dimension.value,
        "dimension_zh": row["dimension_zh"],
        "entry_type": "新增规则" if "白皮书补齐" in (rule.tags or []) else "既有规则",
        "target_column": target_column,
        "trial_total_rows": trial.total_rows,
        "trial_failure_count": trial.failure_count,
        "trial_pass_rate": trial.pass_rate,
        "task_run_status": run.status.value,
        "issue_count": len(run.issue_ids),
        "dashboard_overall_score": dashboard["overall_score"],
        "script_preview_unresolved": preview["unresolved_placeholders"],
        "result": "PASS",
    }


def _render_markdown(result: Dict[str, Any]) -> str:
    lines = [
        "# 白皮书规则库补齐流程测试报告",
        "",
        f"- 白皮书文件：{result['whitepaper']}",
        f"- 测试时间：{result['tested_at']}",
        f"- 白皮书规则数：{result['whitepaper_rule_count']}",
        f"- 覆盖规则数：{result['covered_rule_count']}",
        f"- 新增规则测试数：{result['added_rule_test_count']}",
        f"- 既有规则测试数：{result['existing_rule_test_count']}",
        f"- 测试结论：{result['test_result']}",
        "",
        "## 测试数据说明",
        "",
        "每条规则使用两行样例数据：第一行通过 `__fail_rules` 标记当前规则失败，第二行为通过样例。这样可以在不连接真实外部数据库、接口、权限系统的情况下，逐条验证规则配置、脚本预览、试跑、任务执行、问题生成、看板与报告输出的完整流程。",
        "",
        "## 规则测试明细",
        "",
        "| 序号 | 白皮书规则 | 入库规则 | 维度 | 类型 | 试跑结果 | 异常行 | 任务状态 | 问题数 |",
        "| --- | --- | --- | --- | --- | --- | ---: | --- | ---: |",
    ]
    for index, item in enumerate(result["rule_results"], start=1):
        lines.append(
            "| {index} | {wp} {name} | {rule} | {dimension} | {entry_type} | {res} | {failures} | {status} | {issues} |".format(
                index=index,
                wp=item["whitepaper_rule_id"],
                name=item["display_name"],
                rule=item["rule_id"],
                dimension=item["dimension_zh"],
                entry_type=item["entry_type"],
                res=item["result"],
                failures=item["trial_failure_count"],
                status=item["task_run_status"],
                issues=item["issue_count"],
            )
        )
    if result["failures"]:
        lines.extend(["", "## 失败详情", ""])
        for failure in result["failures"]:
            lines.append(f"- {failure['whitepaper_key']}：{failure['error']}")
    return "\n".join(lines) + "\n"


def run(whitepaper_path: str = "", output_dir: str = "") -> Dict[str, Any]:
    modules = _load_modules()
    rule_library_mod = modules["rule_library"]
    library = rule_library_mod.RuleLibrary(load_builtin=True)
    rules = library.list_rules()
    whitepaper = _find_whitepaper(whitepaper_path)
    whitepaper_rows = _extract_whitepaper_rules(whitepaper)

    rule_results: List[Dict[str, Any]] = []
    failures: List[Dict[str, Any]] = []
    for row in whitepaper_rows:
        key = _whitepaper_rule_key(row)
        try:
            rule = _find_rule_for_whitepaper_row(rules, row)
            rule_results.append(_test_rule(modules, rule, row))
        except Exception as exc:  # noqa: BLE001 - report every rule failure
            failures.append(
                {
                    "whitepaper_key": key,
                    "whitepaper_rule_id": row["whitepaper_rule_id"],
                    "display_name": row["display_name"],
                    "error": str(exc),
                }
            )

    added_rule_count = sum(1 for item in rule_results if item["entry_type"] == "新增规则")
    existing_rule_count = sum(1 for item in rule_results if item["entry_type"] == "既有规则")
    result = {
        "test_result": "PASS" if not failures else "FAIL",
        "tested_at": datetime.now().replace(microsecond=0).isoformat(),
        "whitepaper": str(whitepaper),
        "whitepaper_rule_count": len(whitepaper_rows),
        "covered_rule_count": len(rule_results),
        "added_rule_test_count": added_rule_count,
        "existing_rule_test_count": existing_rule_count,
        "library_total_rules": len(rules),
        "rule_results": rule_results,
        "failures": failures,
    }

    out_dir = Path(output_dir) if output_dir else DATA_QUALITY / "docs"
    out_dir.mkdir(parents=True, exist_ok=True)
    json_path = out_dir / "whitepaper_rule_flow_test_results.json"
    md_path = out_dir / "whitepaper_rule_flow_test_report.md"
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md_path.write_text(_render_markdown(result), encoding="utf-8")
    result["json_report"] = str(json_path)
    result["markdown_report"] = str(md_path)
    return result


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run whitepaper rule flow tests.")
    parser.add_argument("--whitepaper", default="", help="Path to the whitepaper docx.")
    parser.add_argument("--output-dir", default="", help="Directory for JSON and Markdown reports.")
    args = parser.parse_args(list(argv) if argv is not None else None)
    result = run(args.whitepaper, args.output_dir)
    print(
        json.dumps(
            {
                "test_result": result["test_result"],
                "whitepaper_rule_count": result["whitepaper_rule_count"],
                "covered_rule_count": result["covered_rule_count"],
                "added_rule_test_count": result["added_rule_test_count"],
                "existing_rule_test_count": result["existing_rule_test_count"],
                "failures": result["failures"],
                "json_report": result["json_report"],
                "markdown_report": result["markdown_report"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if result["test_result"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
